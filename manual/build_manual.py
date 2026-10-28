"""Build the operator manual from the verified recovery configuration."""
from copy import deepcopy
from pathlib import Path
import math

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Mm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "Delta robot fly tracking user manual.docx"
FONT = Path("C:/Windows/Fonts/arial.ttf")
BOLD = Path("C:/Windows/Fonts/arialbd.ttf")
BLACK = "000000"


def pil_font(size, bold=False):
    return ImageFont.truetype(str(BOLD if bold else FONT), size)


def center_text(draw, box, text, size=31, bold=False):
    font = pil_font(size, bold)
    bounds = draw.multiline_textbbox((0, 0), text, font=font, spacing=7, align="center")
    x = (box[0] + box[2] - (bounds[2] - bounds[0])) / 2
    y = (box[1] + box[3] - (bounds[3] - bounds[1])) / 2 - bounds[1]
    draw.multiline_text((x, y), text, fill="#18232C", font=font, spacing=7, align="center")


def arrow(draw, points, fill="#44677B", width=8):
    draw.line(points, fill=fill, width=width, joint="curve")
    a, b = points[-2], points[-1]
    angle = math.atan2(b[1] - a[1], b[0] - a[0])
    tip = [b, (b[0] - 19 * math.cos(angle - .48), b[1] - 19 * math.sin(angle - .48)),
           (b[0] - 19 * math.cos(angle + .48), b[1] - 19 * math.sin(angle + .48))]
    draw.polygon(tip, fill=fill)


def figures():
    ASSETS.mkdir(exist_ok=True)
    im = Image.new("RGB", (1900, 500), "white")
    d = ImageDraw.Draw(im)
    boxes = [(30, 90, 410, 260), (500, 90, 880, 260), (970, 90, 1350, 260), (1440, 90, 1820, 260)]
    labels = ["Fixed tracking\ncamera", "Find the fly\nand its offset", "Calculate\nX and Y moves", "Delta robot\nmoves the arena"]
    for box, label in zip(boxes, labels):
        d.rounded_rectangle(box, radius=12, fill="#F2F5F7", outline="#97A8B2", width=3)
        center_text(d, box, label, 33)
    for a, b in zip(boxes, boxes[1:]):
        arrow(d, [(a[2] + 10, 175), (b[0] - 15, 175)])
    arrow(d, [(1630, 275), (1630, 365), (220, 365), (220, 275)])
    center_text(d, (450, 385, 1440, 455), "The next image shows the result of the movement", 30)
    im.save(ASSETS / "control-loop.png")

    im = Image.new("RGB", (1600, 480), "white")
    d = ImageDraw.Draw(im)
    d.rectangle((40, 30, 660, 430), fill="#FAFAFA", outline="#80929D", width=3)
    cx, cy = 350, 230
    d.line((cx - 35, cy, cx + 35, cy), fill="#7D8E98", width=3)
    d.line((cx, cy - 35, cx, cy + 35), fill="#7D8E98", width=3)
    d.ellipse((cx + 105, cy - 25, cx + 155, cy + 25), fill="#17242D")
    arrow(d, [(cx + 5, cy - 65), (cx + 125, cy - 65)], width=5)
    d.text((cx + 15, cy - 108), "+40 pixels X", font=pil_font(27), fill="#18232C")
    d.text((cx - 70, cy + 55), "Image center", font=pil_font(25), fill="#18232C")
    center_text(d, (740, 70, 1550, 175), "40 / -100 = -0.4", 55, True)
    center_text(d, (740, 190, 1550, 310), "Default X command\nG01 X-0.40000 Y0.00000", 32)
    center_text(d, (740, 330, 1550, 430), "Controller units and axis signs\nmust be established on the rig", 26)
    im.save(ASSETS / "gain-example.png")


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin, sec.bottom_margin = Mm(18), Mm(17)
sec.left_margin = sec.right_margin = Mm(19)
sec.header_distance = sec.footer_distance = Mm(8)

for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption", "List Bullet"]:
    style = doc.styles[style_name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor.from_string(BLACK)
    style.font.italic = False
    style.font.underline = False
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")

normal = doc.styles["Normal"]
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
for name, size, before, after in [("Title", 28, 0, 7), ("Subtitle", 15, 0, 10),
                                  ("Heading 1", 19, 0, 10), ("Heading 2", 12, 10, 5)]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.font.bold = name != "Subtitle"
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
doc.styles["Caption"].font.size = Pt(9)
doc.styles["Caption"].font.italic = False

header = sec.header.paragraphs[0]
header.text = "Delta robot fly tracking"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(8)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("User manual 1.0  |  Page ").font.size = Pt(8)
for instr in ["PAGE", "NUMPAGES"]:
    if instr == "NUMPAGES":
        footer.add_run(" of ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instr)
    footer._p.append(field)

doc.core_properties.title = "Delta robot fly tracking user manual"
doc.core_properties.subject = "Operating the recovered Bonsai fly arena controller"
doc.core_properties.author = "Laboratory operations"
doc.core_properties.keywords = "Bonsai 2.7.1, delta robot, fly tracking, operator manual"
doc.core_properties.comments = "Recovery edition 28 October 2026"


def p(text="", bold=False, style=None, after=None):
    para = doc.add_paragraph(style=style)
    para.add_run(text).bold = bold
    if after is not None:
        para.paragraph_format.space_after = Pt(after)
    return para


def h(text, level=2):
    return doc.add_paragraph(text, style=f"Heading {level}")


def step(number, lead, text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Mm(5)
    para.paragraph_format.first_line_indent = Mm(-5)
    para.add_run(f"{number}  {lead} ").bold = True
    para.add_run(text)
    return para


def bullet(text):
    return p(text, style="List Bullet")


def code(text):
    para = p(text)
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after = Pt(5)
    for run in para.runs:
        run.font.name = "Consolas"
        run.font.size = Pt(9)
    return para


def table(headers, rows, widths):
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    for col, width in zip(tab.columns, widths):
        col.width = Mm(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        item = OxmlElement("w:" + edge)
        for key, val in {"val": "single", "sz": "5", "color": "D9D9D9"}.items():
            item.set(qn("w:" + key), val)
        borders.append(item)
    tab._tbl.tblPr.append(borders)
    for i, text in enumerate(headers):
        tab.rows[0].cells[i].text = text
    repeat = OxmlElement("w:tblHeader")
    tab.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = tab.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    for ri, row in enumerate(tab.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        for ci, cell in enumerate(row.cells):
            cell.width = Mm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge, value in [("top", "85"), ("bottom", "85"), ("left", "105"), ("right", "105")]:
                item = OxmlElement("w:" + edge)
                item.set(qn("w:w"), value)
                item.set(qn("w:type"), "dxa")
                margins.append(item)
            tcpr.append(margins)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E8EFF4" if ri == 0 else ("F7F8FA" if ri % 2 == 0 else "FFFFFF"))
            tcpr.append(shading)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs:
                    run.font.size = Pt(10)
                    run.bold = ri == 0
    p("", after=2)
    return tab


def picture(file, width, alt, floating=False):
    para = p("", after=3)
    if not floating:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = para.add_run().add_picture(str(ASSETS / file), width=Mm(width))
    shape._inline.docPr.set("descr", alt)
    if floating:
        inline = shape._inline
        anchor = OxmlElement("wp:anchor")
        for key, value in {"distT": "0", "distB": "91440", "distL": "182880", "distR": "0", "simplePos": "0",
                           "relativeHeight": "0", "behindDoc": "0", "locked": "0", "layoutInCell": "1", "allowOverlap": "0"}.items():
            anchor.set(key, value)
        pos = OxmlElement("wp:simplePos"); pos.set("x", "0"); pos.set("y", "0"); anchor.append(pos)
        for axis, ref, value in [("H", "column", "right"), ("V", "paragraph", "0")]:
            position = OxmlElement("wp:position" + axis); position.set("relativeFrom", ref)
            child = OxmlElement("wp:align" if axis == "H" else "wp:posOffset"); child.text = value
            position.append(child); anchor.append(position)
        anchor.append(deepcopy(inline.extent))
        wrap = OxmlElement("wp:wrapSquare"); wrap.set("wrapText", "left"); anchor.append(wrap)
        for tag in ["wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"]:
            element = inline.find(qn(tag))
            if element is not None: anchor.append(deepcopy(element))
        parent = inline.getparent(); parent.remove(inline); parent.append(anchor)
    return para


def page(title):
    doc.add_page_break()
    h(title, 1)


def field(label):
    label_para = p(label, after=0)
    label_para.paragraph_format.keep_with_next = True
    para = p("_" * 68, after=7)
    para.paragraph_format.line_spacing = 1.0


def repository_link():
    para = p("Project repository and downloads  ")
    url = "https://github.com/itskovpa/delta-fly-microscope-recovery"
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), para.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); props.append(underline)
    run.append(props)
    text = OxmlElement("w:t"); text.text = "Open the public repository"; run.append(text)
    link.append(run); para._p.append(link)


figures()

# PAGE 1
p("Delta robot fly tracking", style="Title")
p("User manual", style="Subtitle")
p("Version 1.0  |  Recovery edition  |  28 October 2026", after=12)
p("The system finds a fly in a fixed camera image and moves the arena with a delta robot to keep the fly near the image center. This manual explains the recovered Bonsai software, its controls and the checks needed before physical use.")
p("Start with the offline demo. Physical operation requires a commissioned rig with a recorded camera setup, controller configuration and approved startup procedure.", bold=True)
picture("control-loop.png", 172, "Feedback loop from fixed camera to fly position, movement command and delta robot arena movement")
h("Choose your starting point")
table(["Task", "Where to go"], [
    ("Learn the controls without hardware", "Practice without hardware  |  page 2"),
    ("Install on another Windows computer", "Install the software  |  page 3"),
    ("Reconnect or change the physical setup", "Commission the rig  |  page 4"),
    ("Use an already commissioned setup", "Run and stop an experiment  |  page 5"),
    ("Check settings or resolve a problem", "Controls  |  page 6    Troubleshooting  |  page 7"),
    ("Record the approved setup", "Rig configuration record  |  page 8"),
], [92, 80])
h("Understand the three stop actions")
p("End stops new vision-derived commands. Shift+F5 stops the Bonsai workflow. Neither proves that queued robot movement has stopped. Use the rig's physical emergency stop when the robot must stop immediately.")
p("Starting the real controller can immediately send homing and Z movement commands, before S is pressed. Open it for inspection until commissioning is complete.", bold=True)

# PAGE 2
page("Practice without hardware")
p("Use Open Offline Demo.cmd in the recovery folder. It starts an artificial camera and calculates command strings without opening a camera, serial port or robot connection.")
p("The key sequence below comes from the saved workflow. Confirm it in this demo before physical use; keyboard interaction has not yet been certified.")
picture("offline-workflow.png", 54, "Actual exported offline Bonsai graph with CommandsGain, Tracking and CommandProbe nodes", floating=True)
step(1, "Open the demo.", "Double-click Open Offline Demo.cmd. Use this launcher so the correct Bonsai environment is selected.")
step(2, "View the artificial camera.", "Double-click Tracking to open its group. While the workflow runs, double-click SyntheticFly to open its visualizer. Expect a dark disk at changing positions and occasional fly-absent frames.")
step(3, "View commands.", "Return to the top-level graph and open CommandProbe. The graph at right shows the node names to find. A Command node carries strings between the groups.")
step(4, "Start tracking output.", "Focus the Bonsai editor and press S. Expect strings such as G01 X-0.40000 Y-0.20000. These are displayed only; they are not sent to a robot.")
step(5, "Stop tracking output.", "Press End. The command gate closes while the artificial image stream continues. End does not restart or stop the whole workflow.")
step(6, "Start a new attempt.", "Press Shift+F5, then F5, then S. The saved gate is not designed to reopen with S after End unless the workflow is restarted.")
step(7, "Finish.", "Press Shift+F5 and close the demo. Do not replace its synthetic source with a real camera or add a serial writer.")
h("What the demo demonstrates")
p("The same image-processing and command-building logic is exercised using known image positions. It teaches the graph and control sequence. It does not model robot mechanics, validate camera drivers or establish safe physical gain values.")
p("If a visualizer is closed, double-click its source node again while running. If the demo does not start or the keys have no effect, use the checks on page 7.")

# PAGE 3
page("Install the software")
p("An installed recovery copy is ready on this computer. Use these steps for a new Windows PC or a fresh copy of the recovery folder. Installation and validation can be completed without hardware.")
h("Computer requirements")
p("Use 64-bit Windows with .NET Framework 4.7.2 or later, a .NET 8 or 9 SDK, and Python 3. Internet access is needed for the first restoration. Keep the recovery folder in a writable location and retain its originals, environment and local-packages folders.")
step(1, "Get the recovery folder.", "Copy the complete folder or download the public repository listed on page 8. Do not copy only the .bonsai workflow. The system Bonsai installation is separate and may use a different version.")
step(2, "Restore the environment.", "Open PowerShell in the recovery folder and run this command as one line. The script installs the pinned packages under runtime.")
code(r"powershell.exe -NoProfile -ExecutionPolicy Bypass -File .\Restore-Environment.ps1")
step(3, "Validate without hardware.", "Close the recovery editor, then run the following command. Wait for both April variants to report passed true. Each test expects 80 frames, 8 fly-absent frames and 72 tracking commands.")
code(r"powershell.exe -NoProfile -ExecutionPolicy Bypass -File .\Validate-Offline.ps1")
step(4, "Open the offline demo.", "Use Open Offline Demo.cmd and complete page 2. If restoration or validation reports an error, do not proceed to the physical controller.")
h("Pinned software versions")
table(["Component", "Version"], [
    ("Bonsai executable, Editor and Core", "2.7.1"),
    ("Vision, System, Dsp, Windows.Input and expression scripting", "2.7.0"),
    ("OpenCV.Net / native OpenCV Windows packages", "3.4.1 / 2.4.13"),
    ("Bundled Bonsai.MvCamCtrl / MvCamCtrl.Net", "0.1.0 / 4.0.0"),
], [127, 45])
p("There are 27 runtime packages in the saved environment. Use its launchers rather than opening workflows through a general .bonsai file association. Do not update packages on a working rig without a separate test copy.")
p("A message about rechecking startup after installation is expected on a fresh restore: the script verifies package files and requires a successful second launch to handle a legacy bootstrapper issue.")

# PAGE 4
page("Commission the rig")
p("Commissioning is required before the first physical run and after changing a camera, controller, firmware, geometry or axis calibration. An experienced person must approve the rig procedure and complete the record on page 8.")
p("Keep robot motion inhibited using the approved rig procedure while checking cameras and reviewing startup commands. Use the rig-specific power-on, homing and stop procedures recorded on page 8.", bold=True)
step(1, "Identify the hardware.", "Record robot model, firmware, serial adapter, units, workspace limits, homing direction and physical emergency-stop action. Verify the controller's required communication settings from its documentation.")
step(2, "Check the camera without robot output.", "Use Open Camera Check.cmd. It opens Camera-Tracking-Check.bonsai without starting. Press F5 only when the camera is available. This diagnostic contains no serial writer, homing, movement or recording branch.")
step(3, "Inspect detection.", "Double-click Camera Tracking Check to view CameraCapture, Threshold and the centroid nodes. Select the correct camera index. The fly should be bright in the inverse-threshold mask and be the largest accepted region. The exposed CenteredXYErrorPx output is the X/Y pixel error; (0, 0) is image center. Check fly-absent images too. Stop with Shift+F5.")
step(4, "Review the real controller.", "Use Open Controller for Inspection.cmd. Inspect CreateSerialPort and the startup branches without pressing F5. The old port is COM19, with 115200 baud, 8 data bits, no parity and 2 stop bits. Verify those settings for the actual controller.")
step(5, "Approve startup and movement handling.", "The graph has a separate G28 branch and a sequence containing G90, G01 Z-263 F100, G91 and motion-limit settings. It does not wait for a homing acknowledgement. The saved Z coordinate and limits must be reviewed for this rig before enabling motion.")
step(6, "Establish scale and axis direction.", "Use an approved bounded test procedure with a stationary target. Check one small correction at a time. Confirm that X and Y move the fly image toward the center. Record calibrated divisors, operating bounds and the procedure for clearing queued motion.")
step(7, "Approve routine operation.", "Confirm stable centering, a reliable stop response, logging and any required recording. Train users on page 5. Do not use a live fly to discover axis direction or an unverified movement scale.")
h("Camera types")
p("The tracking branch uses generic CameraCapture index 0. The two Hikrobot recording branches are disabled. If Hikrobot cameras are later enabled, their native MVS driver/runtime must match the rig; installing a NuGet package alone does not install the USB or GigE driver.")

# PAGE 5
page("Run and stop an experiment")
p("Use this procedure only after commissioning is signed off and the approved rig startup procedure is available. It applies to controller/Delta-Controller.bonsai, not the offline demo.")
h("Before starting")
bullet("Use the recovery environment and check the rig record for the approved camera, COM port, thresholds and scaling divisors.")
bullet("Check the arena attachment, available workspace and emergency stop using the approved physical procedure.")
bullet("If camera detection needs checking, use Open Camera Check.cmd first and close that workflow before opening the real controller; two workflows may compete for the same camera.")
bullet("Confirm the output folder is writable. The active log records command strings. The two auxiliary video recorders are disabled in this controller.")
h("Start and observe")
step(1, "Open for inspection.", "Double-click Open Controller for Inspection.cmd. This opens the real workflow without starting it.")
step(2, "Start the workflow.", "Only when the approved physical conditions are met, press F5. Homing and Z movement commands may be sent immediately. S is not a master motion-enable switch.")
step(3, "Confirm readiness.", "Follow the approved homing/startup confirmation procedure. Inspect the tracking image and centroid before tracking output. The image-processing group is unnamed in this controller; identify it by its CameraCapture node.")
step(4, "Start correction.", "Press S when the rig is ready. Observe the image and motion continuously at first. The fly should remain near the image center. If motion diverges, oscillates or approaches a limit, stop using the rig's approved procedure.")
step(5, "Monitor the session.", "Check that the detected region remains the fly and that commands/logging are updating. A valid contour can belong to a reflection or another object; contour validity alone does not prove correct tracking.")
h("Stop and save")
step(1, "Stop new tracking commands.", "Press End. For unexpected or unsafe movement, use the physical emergency stop rather than waiting for software keys.")
step(2, "Stop Bonsai.", "Press Shift+F5. Confirm the robot has stopped and resolve any buffered motion using the approved controller procedure.")
step(3, "Check the output.", "Inspect the timestamped commands file in controller/recordings and note the session settings. This is a command log, not measured robot position, video or proof of executed movement.")
p("For another session, repeat the approved startup procedure. Restarting with F5 can send homing and Z commands again.", bold=True)

# PAGE 6
page("Controls and settings")
table(["Control", "Effect", "Important limit"], [
    ("F5", "Start the workflow", "Real controller can send startup motion immediately"),
    ("Shift+F5", "Stop the workflow", "Does not prove the robot queue is clear"),
    ("S", "Open the tracking-command gate", "Use after rig startup is confirmed"),
    ("End", "Close the tracking-command gate", "Restart workflow before another S attempt"),
    ("Arrows and A / Q", "Stored X, Y and Z jog mappings", "Disabled in the saved controller"),
    ("Mouse and Space", "Stored mouse-control mappings", "Disabled in the saved controller"),
], [30, 66, 76])
h("Where to change a setting")
p("Select a node to inspect its properties. Double-click a group to see its contents. Change physical-operation values only under the commissioning procedure; save a dated settings record.")
table(["Setting", "Location", "Saved April 13 value"], [
    ("Camera selection", "CameraCapture > Index", "0; host dependent"),
    ("Exposure", "CameraCapture > CaptureProperties", "-7; camera-specific meaning"),
    ("Segmentation threshold", "Threshold > ThresholdValue", "163; BinaryInv"),
    ("Minimum contour area", "FindContours > MinArea", "20; tune for magnification"),
    ("X and Y response", "CommandsGain > ScalingFactorX / Y", "-100 for each divisor"),
    ("Robot port", "CreateSerialPort > PortName", "COM19; connection name Delta"),
], [42, 78, 52])
picture("gain-example.png", 168, "Example of a fly forty pixels right of image center producing an X command of minus zero point four with divisor minus one hundred")
p("ScalingFactorX and Y are divisors. A larger absolute divisor gives a smaller response; a smaller absolute divisor gives a larger response. Changing the sign reverses that axis. The physical movement units and safe values are established during commissioning.")

# PAGE 7
page("Troubleshooting")
p("If physical motion is unexpected, use the approved stop procedure before troubleshooting. Use the offline demo or camera-only check to isolate software and detection issues.")
table(["Symptom", "Check and action"], [
    ("Wrong Bonsai version opens", "Close it and use the recovery launchers. They select runtime/Bonsai.exe rather than the system file association."),
    ("Port not found or access denied", "Check the assigned COM port and close other serial programs. Match CreateSerialPort to the rig record; keep its connection name Delta."),
    ("No image or camera busy", "Run only one camera workflow. Check camera power, driver and index with Open Camera Check.cmd. Camera indices may change after USB reconnection."),
    ("Detector follows the wrong object", "Inspect the inverse-threshold mask and largest contour. Check lighting, exposure, threshold and MinArea in camera-only mode."),
    ("No tracking commands after S", "Confirm the workflow is running, the fly is detected and the editor has focus. If End was pressed, use Shift+F5 then F5 before S. On the real rig, repeat approved startup."),
    ("Image is not a member of IplImage", "An old active camera branch expects a wrapped image. The surviving MvCamCtrl node already returns IplImage. Have the branch corrected in a separate working copy; do not enable archived recorders blindly."),
    ("OpenCV or native DLL is missing", "Run restoration and offline validation. On a new PC, check the Microsoft Visual C++ 2015-2022 x64 runtime if native loading still fails."),
    ("No videos are saved", "The two auxiliary recorders are disabled. The active command log is not a video recording. Enable and validate recording separately."),
    ("Commands use decimal commas", "Historical formatting follows Windows culture. A commissioner must enforce decimal-point formatting before physical use. The offline test fixes its own culture."),
    ("Oscillation or movement away from center", "Stop motion. Recheck axis signs, scaling, setpoint and latency with a stationary target under the commissioning procedure."),
], [52, 120])
h("When to hand over to the commissioning lead")
p("Stop routine operation after a firmware, camera, geometry or axis-setting change, a failed stop response, or repeated tracking divergence. Record the error, workflow filename, relevant settings and the most recent safe action. A passed offline test does not clear a physical fault.")

# PAGE 8
page("Rig configuration record")
p("Complete this record during commissioning. Keep it with the rig and update it whenever hardware or approved settings change. Use the editable Word copy or write on a printed copy.")
field("Rig identifier and location")
field("Robot model and firmware")
field("Emergency stop and queue clearing procedure reference")
field("Approved startup and homing procedure reference")
field("COM port and verified protocol")
field("Tracking camera model serial number and index")
field("Image size exposure threshold and minimum area")
field("X and Y divisors verified units signs and operating bounds")
field("Recording requirements and output location")
field("Commissioned by and date")
h("Software reference")
p("Use controller/Delta-Controller.bonsai through the inspection launcher. It derives from the newest saved controller dated 13 April 2023, with the command-log path changed to recordings\\commands.csv. The filename 3cameras in the historical sources does not mean all cameras are enabled.")
p("Bonsai 2.7.1 and the April 10 / April 13 pipelines passed compile-only and synthetic tests. Each synthetic test generated 80 frames, including 8 fly-absent frames, and produced 72 expected commands with maximum command-value error 0. Physical performance and actual keyboard interaction were not certified by those tests.")
p("The ImageSta file-analysis and older plotter/Magnet workflows are separate experiments, not the routine robot controller. Their additional Python and legacy camera dependencies are outside this operator procedure.")
p("Full setup details: README.md. Source evidence: WORKFLOW_AUDIT.md. Test evidence: validation/RESULTS.md.")
repository_link()
p("Manual basis: recovery software snapshot 77f929e and the derived camera-only check. Revision 1.0, 28 October 2026.", after=0)

# The bundled default Word template carries a title border. Remove paragraph
# borders from both content and styles so no decorative rule survives in Word.
for root in [doc._element, doc.styles.element]:
    for border in list(root.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
doc.save(OUT)
print(OUT)
